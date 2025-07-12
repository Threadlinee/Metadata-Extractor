from tkinter import Tk, Label, Button, filedialog, Text, Frame, Scrollbar
from PIL import Image, ExifTags
from PyPDF2 import PdfReader
from docx import Document
from openpyxl import load_workbook
import os
import requests
import threading

VT_API_KEY = '676afc2cfea16d7e0fd420f236f87240636437fe69c0a691a0333ab99edb96c6'
VT_BASE_URL = 'https://www.virustotal.com/api/v3/files/'

def extract_image_metadata(file_path):
    metadata = []
    try:
        image = Image.open(file_path)
        exif_data = image.getexif() if hasattr(image, 'getexif') else None
        if exif_data:
            for tag_id, value in exif_data.items():
                tag = ExifTags.TAGS.get(tag_id, tag_id)
                metadata.append(f"{tag}: {value}")
        else:
            metadata.append("No EXIF metadata found.")
    except Exception as e:
        metadata.append(f"Error reading image metadata: {e}")
    return metadata

def extract_pdf_metadata(file_path):
    metadata = []
    try:
        reader = PdfReader(file_path)
        info = reader.metadata
        if info:
            for key, value in info.items():
                metadata.append(f"{key}: {value}")
        else:
            metadata.append("No PDF metadata found.")
    except Exception as e:
        metadata.append(f"Error reading PDF metadata: {e}")
    return metadata

def extract_docx_metadata(file_path):
    metadata = []
    try:
        doc = Document(file_path)
        core_props = doc.core_properties
        for prop in dir(core_props):
            if not prop.startswith('_') and not callable(getattr(core_props, prop)):
                value = getattr(core_props, prop)
                if value:
                    metadata.append(f"{prop}: {value}")
    except Exception as e:
        metadata.append(f"Error reading Word metadata: {e}")
    return metadata

def extract_xlsx_metadata(file_path):
    metadata = []
    try:
        wb = load_workbook(file_path)
        props = wb.properties
        for prop in dir(props):
            if not prop.startswith('_') and not callable(getattr(props, prop)):
                value = getattr(props, prop)
                if value:
                    metadata.append(f"{prop}: {value}")
    except Exception as e:
        metadata.append(f"Error reading Excel metadata: {e}")
    return metadata

def virustotal_check(file_path, sha256):
    headers = {
        'x-apikey': VT_API_KEY
    }
    vt_result = []
    try:
        resp = requests.get(VT_BASE_URL + sha256, headers=headers)
        if resp.status_code == 200:
            data = resp.json()
            stats = data['data']['attributes']['last_analysis_stats']
            total = sum(stats.values())
            malicious = stats.get('malicious', 0)
            suspicious = stats.get('suspicious', 0)
            harmless = stats.get('harmless', 0)
            undetected = stats.get('undetected', 0)
            vt_result.append(f"VirusTotal: {malicious + suspicious}/{total} engines flagged this file.")
            if malicious + suspicious > 0:
                vt_result.append("⚠️ This file is potentially MALICIOUS!")
            elif harmless > 0 and malicious + suspicious == 0:
                vt_result.append("✅ This file appears CLEAN.")
            else:
                vt_result.append("❓ No conclusive result.")
        elif resp.status_code == 404:
            # Not found, upload file
            vt_result.append("File not found on VirusTotal. Uploading for analysis...")
            with open(file_path, 'rb') as f:
                files = {'file': (os.path.basename(file_path), f)}
                upload_resp = requests.post('https://www.virustotal.com/api/v3/files', headers=headers, files=files)
                if upload_resp.status_code == 200:
                    vt_result.append("File uploaded to VirusTotal. Please check back later for results.")
                else:
                    vt_result.append(f"Error uploading to VirusTotal: {upload_resp.status_code}")
        else:
            vt_result.append(f"VirusTotal API error: {resp.status_code}")
    except Exception as e:
        vt_result.append(f"VirusTotal check failed: {e}")
    return vt_result

def extract_pe_metadata(file_path):
    metadata = []
    try:
        import pefile
        pe = pefile.PE(file_path)
        opt_header = getattr(pe, 'OPTIONAL_HEADER', None)
        if opt_header:
            entry_point = getattr(opt_header, 'AddressOfEntryPoint', None)
            image_base = getattr(opt_header, 'ImageBase', None)
            subsystem = getattr(opt_header, 'Subsystem', None)
            if entry_point is not None:
                metadata.append(f"Entry point: {hex(entry_point)}")
            if image_base is not None:
                metadata.append(f"Image base: {hex(image_base)}")
            if subsystem is not None and hasattr(pefile, 'SUBSYSTEM_TYPE'):
                metadata.append(f"Subsystem: {pefile.SUBSYSTEM_TYPE.get(subsystem, subsystem)}")
        if hasattr(pe, 'sections'):
            metadata.append(f"Sections: {len(pe.sections)}")
            for section in pe.sections:
                name = getattr(section, 'Name', b'').decode(errors='ignore').strip()
                size = getattr(section, 'SizeOfRawData', 'Unknown')
                metadata.append(f"Section: {name} Size: {size}")
        # Add VirusTotal check
        import hashlib
        with open(file_path, 'rb') as f:
            sha256 = hashlib.sha256(f.read()).hexdigest()
        metadata.append(f"SHA256: {sha256}")
        metadata.extend(virustotal_check(file_path, sha256))
    except Exception as e:
        metadata.append(f"Error reading PE metadata: {e}")
    return metadata

def extract_generic_metadata(file_path):
    import hashlib, mimetypes, os
    metadata = []
    try:
        size = os.path.getsize(file_path)
        metadata.append(f"File size: {size} bytes")
        mime, _ = mimetypes.guess_type(file_path)
        metadata.append(f"MIME type: {mime}")
        with open(file_path, 'rb') as f:
            data = f.read()
            md5 = hashlib.md5(data).hexdigest()
            sha256 = hashlib.sha256(data).hexdigest()
        metadata.append(f"MD5: {md5}")
        metadata.append(f"SHA256: {sha256}")
    except Exception as e:
        metadata.append(f"Error reading generic metadata: {e}")
    return metadata

def extract_metadata(file_path):
    ext = file_path.lower().split('.')[-1]
    if ext in ['jpg', 'jpeg', 'png', 'tiff', 'ico']:
        return extract_image_metadata(file_path)
    elif ext == 'pdf':
        return extract_pdf_metadata(file_path)
    elif ext == 'docx':
        return extract_docx_metadata(file_path)
    elif ext == 'xlsx':
        return extract_xlsx_metadata(file_path)
    elif ext == 'exe':
        return extract_pe_metadata(file_path)
    else:
        return extract_generic_metadata(file_path)

def scan_with_animation(file, scan_func):
    output.delete(1.0, "end")
    output.insert("end", f"Scanning {file} ", "filename")
    anim_running = [True]
    result_holder = [[]]
    def animate():
        dots = 0
        def step():
            nonlocal dots
            if not anim_running[0]:
                return
            output.delete("end-2l", "end")
            output.insert("end", "." * (dots % 4) + "\n")
            output.update()
            dots += 1
            output.after(400, step)
        step()
    def do_scan():
        result = scan_func(file)
        # Ensure result_holder[0] is always a list of strings
        if not isinstance(result, list):
            result = [str(result)] if result is not None else []
        result_holder[0] = result
        anim_running[0] = False
        # Show results in main thread
        def show_results():
            output.delete(1.0, "end")
            output.insert("end", f"📄 File: {file}\n", "filename")
            for line in result_holder[0]:
                output.insert("end", f"{line}\n")
            output.insert("end", "\n" + "-"*60 + "\n\n")
        output.after(0, show_results)
    # Start animation and scan in parallel
    threading.Thread(target=do_scan, daemon=True).start()
    animate()

def open_files():
    files = filedialog.askopenfilenames(title="Select Files")
    for file in files:
        scan_with_animation(file, extract_metadata)

# GUI Setup
root = Tk()
root.title("Metadata Extractor Pro")
root.geometry("800x600")
root.config(bg="#1e1e2f")

title = Label(root, text="📂 Metadata Extractor Pro", font=("Helvetica", 20, "bold"), bg="#1e1e2f", fg="cyan")
title.pack(pady=20)

btn_frame = Frame(root, bg="#1e1e2f")
btn_frame.pack()

select_btn = Button(btn_frame, text="Select Files", command=open_files, font=("Helvetica", 14), bg="#444", fg="white", padx=20, pady=10)
select_btn.pack()

output_frame = Frame(root, bg="#1e1e2f")
output_frame.pack(pady=10, fill="both", expand=True)

scrollbar = Scrollbar(output_frame)
scrollbar.pack(side="right", fill="y")

output = Text(output_frame, wrap="word", font=("Courier", 11), yscrollcommand=scrollbar.set, bg="#121219", fg="#00FFAA", insertbackground='white')
output.pack(side="left", fill="both", expand=True)

output.tag_config("filename", foreground="cyan", font=("Courier", 11, "bold"))

scrollbar.config(command=output.yview)

root.mainloop()
